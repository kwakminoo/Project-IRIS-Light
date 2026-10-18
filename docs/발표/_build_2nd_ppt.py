# 2026 오픈소스 개발자대회 2차 발표.
# 구성: docs/발표/2차발표-구성과-방향-검토.md 12장.
# 2·3장은 그 다음 합의(학생·초보 개발자, 통로와 포트)를 반영.
# 화면: 흰 종이, 검정 글자, 슬라이드당 주장 1개. 기능 카탈로그·다크 배경 없음.

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = r"c:\Users\kwakm\Desktop\Cusor-Project\Project-IRIS-Light-main\docs\발표\IRIS-2차발표-1차안.pptx"

PAPER = "F4F1EA"
INK = "1A1814"
MUTED = "3F3A34"
TEAL = "1A4D57"
LINE = "D9D3C8"
CARD = "FFFCF7"
FONT = "Malgun Gothic"

W = 13.333
H = 7.5


def rgb(value):
    return RGBColor.from_string(value)


def style_run(run, size, bold, color, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rpr.find(qn(f"a:{tag}"))
        if el is None:
            el = rpr.makeelement(qn(f"a:{tag}"))
            rpr.append(el)
        el.set("typeface", font)


def add_text(slide, x, y, w, h, text, size=20, bold=False, color=INK, align="left", anchor="top"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf._txBody.bodyPr.set("anchor", {"top": "t", "middle": "ctr", "bottom": "b"}[anchor])
    align_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align_map[align]
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.text = line
        style_run(p.runs[0], size, bold, color)
    return box


def add_rect(slide, x, y, w, h, fill, line=None, line_pt=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_pt)
    shape.shadow.inherit = False
    return shape


def card(slide, x, y, w, h):
    add_rect(slide, x, y, w, h, CARD, LINE, 1.25)
    add_rect(slide, x, y, w, 0.06, TEAL)


def paint(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(PAPER)


def footer(slide, num, total=12):
    add_rect(slide, 0.7, 6.92, 11.93, 0.01, LINE)
    add_text(slide, 0.7, 7.02, 8.5, 0.32, "IRIS  ·  Desktop AI Agent Runtime", 14, False, MUTED)
    add_text(slide, 10.4, 7.02, 2.23, 0.32, f"{num:02d}  /  {total:02d}", 14, False, MUTED, align="right")


def header(slide, kicker, title, subtitle):
    add_text(slide, 0.7, 0.32, 10.5, 0.32, kicker, 16, True, TEAL)
    add_text(slide, 0.7, 0.68, 11.9, 0.62, title, 34, True, INK)
    add_text(slide, 0.7, 1.36, 11.9, 0.7, subtitle, 20, False, MUTED)


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def new(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    paint(slide)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "IRIS — Open Source Desktop AI Agent Runtime"
    prs.core_properties.subject = "2026 오픈소스 개발자대회 2차 발표"
    prs.core_properties.category = "OSS contest"

    cover(prs)
    audience(prs)
    problem(prs)
    boundary(prs)
    reuse(prs)
    core_ext(prs)
    docs(prs)
    license_slide(prs)
    student(prs)
    demo(prs)
    verify(prs)
    close(prs)

    prs.save(OUT)
    print(OUT)


def cover(prs):
    s = new(prs)
    add_rect(s, 0, 0, W, 0.08, TEAL)
    add_text(s, 0.75, 0.48, 8, 0.36, "2026 오픈소스 개발자대회", 18, True, TEAL)
    add_text(s, 0.72, 1.22, 7.4, 1.35, "IRIS", 78, True, INK)
    add_rect(s, 0.78, 2.68, 2.1, 0.045, TEAL)
    add_text(s, 0.75, 2.88, 7.2, 0.55, "Desktop AI Agent Runtime", 28, True, INK)
    add_text(s, 0.75, 3.5, 6.8, 0.9, "개발자가 가져다 쓰는\n데스크톱 오픈소스 런타임", 22, False, MUTED)

    add_text(s, 0.75, 5.15, 6.6, 1.35, "GPL-3.0-or-later\n학생 부문\ngithub.com/kwakminoo/Project-IRIS-Light", 18, False, INK)

    card(s, 8.15, 1.45, 4.4, 4.55)
    add_text(s, 8.42, 1.7, 3.9, 0.36, "맡기는 경계", 16, True, TEAL)
    add_text(s, 8.42, 2.2, 3.9, 0.85, "IRIS", 28, True, INK)
    add_text(s, 8.42, 2.95, 3.9, 0.7, "세션 · 권한\n스트리밍 UI", 18, False, MUTED)
    add_rect(s, 8.42, 3.75, 3.85, 0.012, LINE)
    add_text(s, 8.42, 3.95, 3.9, 0.7, "Ollama\n추론", 18, False, INK)
    add_rect(s, 8.42, 4.75, 3.85, 0.012, LINE)
    add_text(s, 8.42, 4.95, 3.9, 0.7, "Hermes\n도구 · 스킬", 18, False, INK)
    footer(s, 1)
    notes(s, "세션과 화면은 IRIS가 맡고, 추론은 Ollama, 도구는 Hermes가 맡습니다. 저장소는 github.com/kwakminoo/Project-IRIS-Light 입니다.")


def audience(prs):
    s = new(prs)
    header(s, "청중", "쓰는 사람은 개발자입니다", "대학생과 처음 만드는 개발자가 가져가 맞춥니다.")
    items = [
        ("01", "포크", "저장소를\n가져갑니다"),
        ("02", "확장", "스킬, MCP, 화면을\n얹습니다"),
        ("03", "성취", "자기 에이전트에\n맞춥니다"),
    ]
    x = 0.7
    for num, title, body in items:
        card(s, x, 2.55, 3.7, 3.55)
        add_text(s, x + 0.28, 2.85, 3.1, 0.4, num, 16, True, TEAL)
        add_text(s, x + 0.28, 3.4, 3.1, 0.7, title, 32, True, INK)
        add_text(s, x + 0.28, 4.3, 3.1, 1.2, body, 22, False, MUTED)
        x += 4.05
    footer(s, 2)
    notes(s, "대학생과 처음 만드는 개발자입니다. 가져가서 스킬, MCP, 화면을 자기 에이전트에 맞춥니다.")


def problem(prs):
    s = new(prs)
    header(s, "문제", "연결이 매번 끊깁니다", "로컬 에이전트는 통로에 흩어져 있습니다.")
    labels = ["터미널", "텔레그램", "디스코드"]
    x = 0.7
    for label in labels:
        card(s, x, 2.45, 3.7, 2.35)
        add_text(s, x + 0.2, 2.7, 3.3, 1.85, label, 32, True, INK, align="center", anchor="middle")
        x += 4.05
    add_text(s, 0.7, 5.15, 11.9, 1.15, "모델 · 도구 · 화면은 포트로만 이어집니다.\n런타임과 게이트웨이가 어긋나면 그 연결이 끊깁니다.", 22, False, INK)
    footer(s, 3)
    notes(s, "로컬 에이전트는 터미널, 텔레그램, 디스코드에서 다루고, 모델과 도구와 화면은 포트로만 이어집니다. 런타임과 게이트웨이가 어긋나면 그 연결이 매번 끊깁니다.")


def boundary(prs):
    s = new(prs)
    header(s, "경계", "IRIS가 고정하는 자리", "세션과 화면은 IRIS가, 추론과 도구는 바깥 런타임이 맡습니다.")
    card(s, 2.15, 2.25, 9.05, 1.45)
    add_text(s, 2.45, 2.4, 8.5, 0.32, "데스크톱 런타임", 16, True, TEAL)
    add_text(s, 2.45, 2.78, 8.5, 0.65, "IRIS", 32, True, INK)
    add_text(s, 5.3, 2.82, 5.5, 0.62, "세션 · 권한\n스트리밍 UI · 시작", 18, False, MUTED)
    # 위임 트리. 인쇄에서 선만 남는다.
    add_rect(s, 6.655, 3.7, 0.02, 0.32, TEAL)
    add_rect(s, 3.55, 4.0, 6.23, 0.02, TEAL)
    add_rect(s, 3.55, 4.0, 0.02, 0.22, TEAL)
    add_rect(s, 9.76, 4.0, 0.02, 0.22, TEAL)
    card(s, 0.7, 4.22, 5.75, 1.7)
    add_text(s, 1.0, 4.4, 5.2, 0.3, "추론", 16, True, TEAL)
    add_text(s, 1.0, 4.75, 5.2, 0.5, "Ollama", 28, True, INK)
    add_text(s, 1.0, 5.32, 5.2, 0.38, ":11434", 18, False, MUTED)
    card(s, 6.88, 4.22, 5.75, 1.7)
    add_text(s, 7.18, 4.4, 5.2, 0.3, "도구 · 스킬", 16, True, TEAL)
    add_text(s, 7.18, 4.75, 5.2, 0.5, "Hermes", 28, True, INK)
    add_text(s, 7.18, 5.32, 5.2, 0.38, ":8642", 18, False, MUTED)
    add_text(s, 0.7, 6.25, 11.9, 0.45, "웹검색, 셸, 파일 입출력은 다시 만들지 않습니다.", 18, False, MUTED)
    footer(s, 4)
    notes(s, "IRIS는 세션, 권한, 스트리밍 화면, 시작 프로토콜을 고정합니다. 검색과 셸과 파일은 다시 만들지 않습니다. 모델은 Ollama 11434, 도구는 Hermes 8642입니다.")


def reuse(prs):
    s = new(prs)
    header(s, "오픈소스", "남이 가져가는 길", "소스 공개만으로 끝나지 않게, 기여 경로를 열어 두었습니다.")
    rows = [
        ("01", "CONTRIBUTING.md", "기여 절차가 문서에 있습니다."),
        ("02", "스킬과 MCP", "integrations/hermes-skills/  ·  코어 diff 0"),
        ("03", "Good first issue", "첫 이슈부터 시작할 자리가 있습니다."),
    ]
    y = 2.35
    for num, title, body in rows:
        card(s, 0.7, y, 11.93, 1.28)
        add_text(s, 1.0, y + 0.28, 1.1, 0.7, num, 22, True, TEAL, anchor="middle")
        add_text(s, 2.3, y + 0.18, 9.8, 0.48, title, 24, True, INK)
        add_text(s, 2.3, y + 0.68, 9.8, 0.4, body, 18, False, MUTED)
        y += 1.45
    footer(s, 5)
    notes(s, "기여 문서는 CONTRIBUTING.md입니다. 스킬은 integrations/hermes-skills/에 넣고 코어 diff는 0으로 둡니다. 첫 이슈는 Good first issue입니다.")


def core_ext(prs):
    s = new(prs)
    header(s, "지속", "코어는 고정합니다", "기여자는 확장부터 넣고, 코어 변경은 최소화합니다.")
    card(s, 0.7, 2.4, 5.75, 3.7)
    add_text(s, 1.0, 2.7, 5.2, 0.35, "CORE", 16, True, TEAL)
    add_text(s, 1.0, 3.2, 5.2, 1.5, "세션\n게이트웨이\n권한", 28, True, INK)
    add_text(s, 1.0, 5.15, 5.2, 0.6, "버그와 보안만 수정합니다.", 18, False, MUTED)
    card(s, 6.88, 2.4, 5.75, 3.7)
    add_text(s, 7.18, 2.7, 5.2, 0.35, "EXTENSION", 16, True, TEAL)
    add_text(s, 7.18, 3.2, 5.2, 1.5, "스킬\nMCP\nIDE", 28, True, INK)
    add_text(s, 7.18, 5.15, 5.2, 0.6, "여기서부터 추가합니다.", 18, False, MUTED)
    add_text(s, 0.7, 6.28, 11.9, 0.42, "회귀는 scripts/run_core_smoke.ps1 한 번입니다.", 18, False, MUTED)
    footer(s, 6)
    notes(s, "코어는 세션, 게이트웨이, 권한입니다. 확장은 스킬, MCP, IDE입니다. 회귀는 scripts/run_core_smoke.ps1 한 번입니다.")


def docs(prs):
    s = new(prs)
    header(s, "문서", "사람과 AI가 같이 읽습니다", "맥락이 보이게 고쳐, 다음 기여자와 에이전트가 이어서 고칩니다.")
    card(s, 0.7, 2.45, 5.75, 3.6)
    add_text(s, 1.0, 2.75, 5.2, 0.4, "사람", 16, True, TEAL)
    add_text(s, 1.0, 3.35, 5.2, 1.6, "domain.md\nIA.md", 32, True, INK)
    add_text(s, 1.0, 5.2, 5.2, 0.5, "구조와 화면의 계약", 18, False, MUTED)
    card(s, 6.88, 2.45, 5.75, 3.6)
    add_text(s, 7.18, 2.75, 5.2, 0.4, "AI", 16, True, TEAL)
    add_text(s, 7.18, 3.55, 5.2, 1.15, "agent-readable-docs.md", 20, True, INK, anchor="middle")
    add_text(s, 7.18, 5.2, 5.2, 0.5, "같은 헤더 계약", 18, False, MUTED)
    footer(s, 7)
    notes(s, "구조 계약은 docs/domain.md와 docs/ia/IA.md입니다. 에이전트용 헤더 계약은 docs/guides/agent-readable-docs.md입니다.")


def license_slide(prs):
    s = new(prs)
    header(s, "라이선스", "GPL인 이유", "UI가 PyQt6라, 결합 저작물은 GPL-3.0-or-later입니다.")
    steps = ["PyQt6", "결합 저작물", "GPL-3.0-or-later"]
    x = 0.7
    for i, title in enumerate(steps):
        card(s, x, 2.55, 3.35, 1.85)
        size = 22 if i == 2 else 28
        add_text(s, x + 0.15, 2.75, 3.05, 1.4, title, size, True, INK, align="center", anchor="middle")
        if i < 2:
            add_text(s, x + 3.32, 3.15, 0.7, 0.55, "→", 26, True, TEAL, align="center", anchor="middle")
        x += 4.05
    add_text(
        s, 0.7, 5.05, 11.9, 1.3,
        "절차는 ui-toolkit-migration.md에 있습니다. 가이드이고, 전환은 아직 하지 않았습니다.",
        20, False, INK,
    )
    footer(s, 8)
    notes(s, "UI가 PyQt6라 결합 저작물은 GPL-3.0-or-later입니다. 다른 UI로 가는 절차는 docs/guides/ui-toolkit-migration.md에 있습니다. 툴킷 전환은 아직 하지 않았습니다.")


def student(prs):
    s = new(prs)
    header(s, "재현", "일반 PC에서 재현됩니다", "재현 조건입니다. 학생 팀의 일반 PC입니다.")
    rows = [
        ("OS", "Windows 10 / 11"),
        ("RAM", "최소 8GB    권장 16GB"),
        ("GPU", "필요 없음. 클라우드 모델이 기본"),
    ]
    y = 2.45
    for key, val in rows:
        card(s, 0.7, y, 11.93, 1.15)
        add_text(s, 1.0, y + 0.28, 2.2, 0.6, key, 22, True, TEAL, anchor="middle")
        add_text(s, 3.4, y + 0.28, 8.8, 0.6, val, 26, True, INK, anchor="middle")
        y += 1.32
    footer(s, 9)
    notes(s, "학생 팀입니다. Windows 10과 11, 메모리 최소 8GB, 권장 16GB, GPU는 필요 없습니다. 클라우드 모델이 기본입니다.")


def demo(prs):
    s = new(prs)
    header(s, "시연", "클립을 재생만 합니다", "발표자는 마우스를 잡지 않습니다.")
    items = [
        ("01", "설치", "30초"),
        ("02", "입력에서 응답", "30초"),
        ("03", "IDE 동반", "30초"),
    ]
    x = 0.7
    for num, title, dur in items:
        card(s, x, 2.45, 3.7, 3.6)
        add_text(s, x + 0.28, 2.75, 3.15, 0.4, num, 16, True, TEAL)
        add_text(s, x + 0.28, 3.35, 3.15, 1.3, title, 28, True, INK)
        add_text(s, x + 0.28, 4.9, 3.15, 0.6, dur, 22, False, MUTED)
        x += 4.05
    footer(s, 10)
    notes(s, "설치, 입력에서 응답, IDE 동반. 각 30초. 마우스는 잡지 않습니다.")


def verify(prs):
    s = new(prs)
    header(s, "검증", "두 가지는 감점 없이", "기능과 라이선스는 발표 점수를 깎는 자리가 아니게 맞춥니다.")
    card(s, 0.7, 2.45, 5.75, 3.15)
    add_text(s, 1.0, 2.75, 5.2, 0.4, "기능", 16, True, TEAL)
    add_text(s, 1.0, 3.3, 5.2, 1.5, "설치\n입력 · 통신 · 응답", 28, True, INK)
    add_text(s, 1.0, 4.95, 5.2, 0.4, "기능테스트-시나리오서.md", 18, False, MUTED)
    card(s, 6.88, 2.45, 5.75, 3.15)
    add_text(s, 7.18, 2.75, 5.2, 0.4, "라이선스", 16, True, TEAL)
    add_text(s, 7.18, 3.3, 5.2, 1.5, "내부 스캔\n충돌 0", 28, True, INK)
    add_text(s, 7.18, 4.95, 5.2, 0.4, "고지와 LICENSE 동기화", 18, False, MUTED)
    add_text(s, 0.7, 5.85, 11.9, 0.7, "대회 온라인 검증 툴 리포트는 검증 기간에 첨부합니다.", 18, False, MUTED)
    footer(s, 11)
    notes(s, "기능은 설치, 입력, 통신, 응답이 기능테스트-시나리오서.md에 있습니다. 라이선스 내부 스캔은 충돌 0입니다. 대회 온라인 툴 리포트는 검증 기간에 붙입니다. 붙기 전에는 끝났다고 말하지 않습니다.")


def close(prs):
    s = new(prs)
    header(s, "맺음", "이 런타임 위에 만듭니다", "가져간 뒤에 만들 수 있는 것, 두 가지입니다.")
    card(s, 0.7, 2.45, 5.75, 2.55)
    add_text(s, 1.0, 2.75, 5.2, 0.4, "01", 16, True, TEAL)
    add_text(s, 1.0, 3.3, 5.2, 1.2, "교육용 에이전트", 28, True, INK)
    card(s, 6.88, 2.45, 5.75, 2.55)
    add_text(s, 7.18, 2.75, 5.2, 0.4, "02", 16, True, TEAL)
    add_text(s, 7.18, 3.3, 5.2, 1.2, "IDE 동반 코딩", 28, True, INK)
    add_text(s, 0.7, 5.35, 11.9, 0.9, "github.com/kwakminoo/Project-IRIS-Light", 26, True, INK)
    footer(s, 12)
    notes(s, "교육용 에이전트와 IDE 동반 코딩을 만들 수 있습니다. 주소만 남기고 말을 멈춥니다.")


if __name__ == "__main__":
    build()
